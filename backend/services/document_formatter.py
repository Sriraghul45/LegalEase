from __future__ import annotations

import html
import re
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

ASSET_DIR = Path(__file__).resolve().parents[2] / "assets"
LOGO_PATH = ASSET_DIR / "legal_ease_logo.png"


def sanitize_text(text: str) -> str:
    """Normalize common typography while keeping useful legal punctuation."""
    replacements = {
        "“": '"',
        "”": '"',
        "‘": "'",
        "’": "'",
        "–": "-",
        "—": "-",
        "…": "...",
        "\u00a0": " ",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return re.sub(r"\r\n?", "\n", text).strip()


def _is_heading(line: str) -> bool:
    clean = line.strip()
    if not clean:
        return False
    return bool(re.match(r"^(?:\d+(?:\.\d+)*[.)]?\s+|[A-Z][A-Z\s&:/-]{4,})", clean))


def format_txt(text: str) -> bytes:
    return sanitize_text(text).encode("utf-8")


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"].font.size = Pt(11)

    if LOGO_PATH.exists():
        logo = doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.add_run().add_picture(str(LOGO_PATH), width=Inches(2.8))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    notice = doc.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    nrun = notice.add_run("DRAFT - FOR REVIEW")
    nrun.bold = True
    nrun.font.size = Pt(10)

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.upper() in {"DRAFT - FOR REVIEW", "DRAFT — FOR REVIEW"}:
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        if _is_heading(line):
            r = p.add_run(line)
            r.bold = True
        else:
            p.add_run(line)

    term_items = [item.strip() for item in terms.split(";") if item.strip()]
    if term_items:
        doc.add_paragraph()
        heading = doc.add_paragraph()
        rr = heading.add_run("Key Terms / Conditions")
        rr.bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "#"
        table.rows[0].cells[1].text = "Condition"
        for idx, item in enumerate(term_items, 1):
            cells = table.add_row().cells
            cells[0].text = str(idx)
            cells[1].text = item

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase - AI-generated draft. Review before use.").font.size = Pt(8)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


class _LegalEasePDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = doc_type
        self.set_auto_page_break(auto=True, margin=18)
        self.set_margins(18, 18, 18)

    def header(self):
        self.set_font("Helvetica", "B", 9)
        self.cell(0, 7, "LegalEase", align="C")
        self.ln(4)

    def footer(self):
        self.set_y(-14)
        self.set_font("Helvetica", "I", 7)
        self.cell(0, 5, "LegalEase - AI-generated draft. Review before use.", align="C")


def format_pdf(text: str, doc_type: str) -> bytes:
    pdf = _LegalEasePDF(doc_type)
    pdf.add_page()
    if LOGO_PATH.exists():
        pdf.image(str(LOGO_PATH), x=70, y=25, w=70)
        pdf.set_y(68)
    pdf.set_x(pdf.l_margin)
    content_width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(content_width, 9, doc_type.upper(), align="C")
    pdf.set_font("Helvetica", "B", 9)
    pdf.multi_cell(content_width, 6, "DRAFT - FOR REVIEW", align="C")
    pdf.ln(3)

    for raw_line in sanitize_text(text).splitlines():
        line = raw_line.strip()
        if not line or line.upper() in {"DRAFT - FOR REVIEW", "DRAFT — FOR REVIEW"}:
            continue
        if _is_heading(line):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(content_width, 6, line)
        else:
            pdf.set_font("Helvetica", "", 10)
            pdf.multi_cell(content_width, 5.5, line)
        pdf.ln(1)

    # fpdf2 returns bytearray in some versions and bytes in others.
    return bytes(pdf.output())


def format_html_preview(text: str) -> str:
    """Convert plain generated text into safe HTML for Streamlit rendering."""
    safe = html.escape(sanitize_text(text))
    blocks = []
    for paragraph in safe.split("\n\n"):
        if not paragraph.strip():
            continue
        blocks.append(f"<p>{paragraph.replace(chr(10), '<br>')}</p>")
    return "".join(blocks)
